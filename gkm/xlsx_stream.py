"""Bound shared-string memory for Excel-authored large XLSX files.

openpyxl read_only streams rows but normally keeps the whole sharedStrings.xml
in RAM. This pinned-3.1.5 adapter spills that table to a temporary SQLite file.
The cache is local, never logged, and deleted on close.
"""

import sqlite3
from tempfile import TemporaryDirectory
from openpyxl.reader.excel import ExcelReader
from openpyxl.cell.text import Text
from openpyxl.xml.functions import iterparse
from openpyxl.xml.constants import SHARED_STRINGS, SHEET_MAIN_NS


class StringTable:
    def __init__(self):
        self.temp = TemporaryDirectory(prefix="gkm-xlsx-strings-")
        self.connection = sqlite3.connect(self.temp.name + "/strings.sqlite3")
        self.connection.execute(
            "CREATE TABLE strings(id INTEGER PRIMARY KEY,value TEXT NOT NULL)"
        )
        self.count = 0

    def load(self, source):
        batch = []
        for _, node in iterparse(source):
            if node.tag == "{" + SHEET_MAIN_NS + "}si":
                value = Text.from_tree(node).content.replace("x005F_", "")
                node.clear()
                batch.append((self.count, value))
                self.count += 1
                if len(batch) >= 1000:
                    self.connection.executemany(
                        "INSERT INTO strings VALUES(?,?)", batch
                    )
                    batch = []
        if batch:
            self.connection.executemany("INSERT INTO strings VALUES(?,?)", batch)
        self.connection.commit()

    def __getitem__(self, index):
        row = self.connection.execute(
            "SELECT value FROM strings WHERE id=?", (index,)
        ).fetchone()
        if row is None:
            raise ValueError("XLSX 共享字符串索引损坏")
        return row[0]

    def close(self):
        self.connection.close()
        self.temp.cleanup()


class Reader(ExcelReader):
    disk_strings = None

    def read_strings(self):
        part = self.package.find(SHARED_STRINGS)
        if part is not None:
            table = StringTable()
            self.disk_strings = table
            with self.archive.open(part.PartName[1:]) as source:
                table.load(source)
            self.shared_strings = table


class Workbook:
    def __init__(self, path):
        self.reader = Reader(path, read_only=True, data_only=False, keep_links=False)
        try:
            self.reader.read()
        except BaseException:
            self.reader.archive.close()
            if self.reader.disk_strings:
                self.reader.disk_strings.close()
            raise
        self.wb = self.reader.wb

    def __getitem__(self, name):
        return self.wb[name]

    @property
    def sheetnames(self):
        return self.wb.sheetnames

    def close(self):
        self.wb.close()
        if self.reader.disk_strings:
            self.reader.disk_strings.close()


def load_workbook(path, **options):
    return Workbook(path)
