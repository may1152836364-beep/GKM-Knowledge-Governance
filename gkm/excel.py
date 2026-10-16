import datetime as dt
import hashlib
import json
import re
import time
from collections import Counter
from pathlib import Path
from openpyxl import Workbook
from .xlsx_stream import load_workbook
from .storage import dumps, uid

ALIASES = {
    "article_id": ["Article ID(文章ID)", "文章ID", "article_id"],
    "title": ["title(文章标题)", "文章标题", "title"],
    "raw_content": ["content(文章正文)", "文章正文", "content", "正文"],
    "language": ["Language(语言)", "语言", "language"],
    "category": ["Category(业务分类)", "业务分类", "category"],
    "release_status": ["Release Status(发布状态)", "发布状态", "release_status"],
    "updated_at": ["Last Edit Time(最后编辑时间)", "最后编辑时间", "updated_at"],
}


def text(cell):
    if cell.value is None:
        return ""
    value = cell.value
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()
    if isinstance(value, (int, float)) and re.fullmatch(
        r"0+", cell.number_format or ""
    ):
        return str(int(value)).zfill(len(cell.number_format))
    return str(value)


def clean(raw):
    # Conservative, reversible transform; never strip business words or HTML tags.
    return raw.replace("\r\n", "\n").replace("\r", "\n")


def parse_status(title, body):
    meaningful = re.sub(r"https?://\S+|<[^>]+>|\s+", "", body)
    if not title.strip():
        return "missing_title"
    if not body.strip():
        return "missing_body"
    if len(meaningful) < 30:
        return "insufficient_content"
    return "ready"


def preview(path, sheet=None):
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        rows = ws.iter_rows()
        headers = [text(c) for c in next(rows)]
        if len(headers) != len(set(headers)) or any(not h for h in headers):
            raise ValueError("列名重复或为空，请先在 Excel 中明确列名，防止元数据丢失")
        mapping = {
            key: next(
                (h for h in headers if h.casefold() in [a.casefold() for a in values]),
                "",
            )
            for key, values in ALIASES.items()
        }
        samples = []
        for row in rows:
            samples.append([text(c) for c in row])
            if len(samples) == 20:
                break
        return {
            "sheets": wb.sheetnames,
            "sheet": ws.title,
            "headers": headers,
            "mapping": mapping,
            "rows": samples,
            "row_estimate": max(0, ws.max_row - 1)
            if ws.max_row is not None
            else "未知（流式导入时精确统计）",
        }
    finally:
        wb.close()


def import_xlsx(
    store, path, mapping, sheet=None, source_scope="default", progress=None
):
    if any(
        not mapping.get(k) for k in ("article_id", "title", "raw_content", "language")
    ):
        raise ValueError(
            "必须映射 ID、标题、正文、语言；缺 ID 的行将隔离，不生成可能错误合并的 ID"
        )
    batch = uid()
    report = {
        "batch_id": batch,
        "total": 0,
        "inserted": 0,
        "updated": 0,
        "duplicates": 0,
        "skipped": 0,
        "errors": [],
        "warnings": [],
        "languages": {},
        "statuses": {},
    }
    languages, statuses = Counter(), Counter()
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        iterator = ws.iter_rows()
        headers = [text(c) for c in next(iterator)]
        if len(set(headers)) != len(headers) or any(not h for h in headers):
            raise ValueError("列名重复或为空")
        for column in mapping.values():
            if column and column not in headers:
                raise ValueError("映射列不存在：" + column)
        with store.connect() as c:
            c.execute(
                "INSERT INTO import_batches VALUES(?,?,?,?)",
                (batch, time.time(), "{}", dumps(mapping)),
            )
            for number, row in enumerate(iterator, 2):
                report["total"] += 1
                values = {h: text(cell) for h, cell in zip(headers, row)}

                def get(k):
                    return values.get(mapping.get(k), "")

                article_id, language = (
                    get("article_id").strip(),
                    get("language").strip(),
                )
                if (
                    not article_id
                    or not language
                    or any(cell.data_type == "f" for cell in row)
                ):
                    report["skipped"] += 1
                    report["errors"].append(
                        {
                            "row": number,
                            "reason": "缺少 ID/语言或包含公式；请核对并以值重新导入",
                        }
                    )
                    continue
                title, raw = get("title"), get("raw_content")
                cleaned = clean(raw)
                status = parse_status(title, cleaned)
                if len(raw) >= 32767:
                    report["warnings"].append(
                        {
                            "row": number,
                            "reason": "正文达到 XLSX 单元格长度上限，需核对源系统是否截断",
                        }
                    )
                if language.lower() not in (
                    "zh",
                    "zh-cn",
                    "en",
                    "en-us",
                    "中文",
                    "英文",
                    "chinese",
                    "english",
                ):
                    report["warnings"].append(
                        {"row": number, "reason": "语言值未识别，保留原值：" + language}
                    )
                languages[language] += 1
                statuses[get("release_status")] += 1
                digest = hashlib.sha256(
                    dumps([title, cleaned, language, values]).encode()
                ).hexdigest()
                existing = c.execute(
                    "SELECT id,current_version FROM articles WHERE article_id=? AND language=? AND source_scope=?",
                    (article_id, language, source_scope),
                ).fetchone()
                if (
                    existing
                    and c.execute(
                        "SELECT 1 FROM article_versions WHERE id=? AND content_hash=?",
                        (existing["current_version"], digest),
                    ).fetchone()
                ):
                    report["duplicates"] += 1
                    continue
                if existing:
                    fk = existing["id"]
                    report["updated"] += 1
                else:
                    fk = c.execute(
                        "INSERT INTO articles(article_id,language,source_scope,category,release_status) VALUES(?,?,?,?,?)",
                        (
                            article_id,
                            language,
                            source_scope,
                            get("category"),
                            get("release_status"),
                        ),
                    ).lastrowid
                    report["inserted"] += 1
                version = c.execute(
                    """INSERT INTO article_versions(article_fk,title,raw_content,clean_content,content_hash,metadata_json,parse_status,source_updated_at,import_batch_id,source_file,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        fk,
                        title,
                        raw,
                        cleaned,
                        digest,
                        dumps(values),
                        status,
                        get("updated_at"),
                        batch,
                        Path(path).name,
                        time.time(),
                    ),
                ).lastrowid
                c.execute(
                    "UPDATE articles SET current_version=?,category=?,release_status=? WHERE id=?",
                    (version, get("category"), get("release_status"), fk),
                )
                if report["total"] % 1000 == 0:
                    c.commit()
                    if progress:
                        progress(report["total"])
            report["languages"], report["statuses"] = dict(languages), dict(statuses)
            c.execute(
                "UPDATE import_batches SET report=? WHERE id=?", (dumps(report), batch)
            )
    except BaseException:
        # Partial batches are explicitly marked; reimport is idempotent.
        with store.connect() as c:
            c.execute(
                "UPDATE import_batches SET report=? WHERE id=?",
                (dumps({**report, "interrupted": True}), batch),
            )
        raise
    finally:
        wb.close()
    store.audit(
        "import",
        batch,
        f"总行数 {report['total']}，新增 {report['inserted']}，更新 {report['updated']}，重复 {report['duplicates']}，隔离 {report['skipped']}",
    )
    return report


def safe_cell(value):
    value = str(value or "")
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def export_results(store, path, job=None):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("治理清单")
    ws.append(
        [
            "文章ID",
            "标题",
            "语言",
            "发布状态",
            "有效性",
            "范围",
            "建议",
            "证据",
            "中文原因",
            "人工结论",
            "复核理由",
            "复核人",
            "任务",
            "导入批次",
            "文章版本",
            "提示词版本",
        ]
    )
    with store.connect() as c:
        query = """SELECT a.article_id,a.language,a.release_status,v.title,v.id version,v.import_batch_id,r.result,j.id job,j.snapshot,
        rv.recommendation review,rv.reason review_reason,rv.reviewer FROM analysis_results r JOIN job_items i ON i.id=r.job_item_fk JOIN analysis_jobs j ON j.id=i.job_fk JOIN article_versions v ON v.id=i.article_version_fk JOIN articles a ON a.id=v.article_fk LEFT JOIN reviews rv ON rv.id=(SELECT MAX(id) FROM reviews WHERE result_fk=r.id) WHERE (? IS NULL OR j.id=?)"""
        for row in c.execute(query, (job, job)):
            result = json.loads(row["result"])
            values = [
                row["article_id"],
                row["title"],
                row["language"],
                row["release_status"],
                result["validity_status"],
                result["affected_scope"],
                result["recommendation"],
                dumps(result["evidence"]),
                result["reason_zh"],
                row["review"],
                row["review_reason"],
                row["reviewer"],
                row["job"],
                row["import_batch_id"],
                row["version"],
                json.loads(row["snapshot"])["prompt"]["id"],
            ]
            ws.append([safe_cell(v) for v in values])
    wb.save(path)
    store.audit("export", "", "已导出治理清单，不含密钥")


def set_cleaning(store, article_pk, mode):
    if mode not in ("preserve-v1", "raw-v1", "html-readable-v1"):
        raise ValueError("未知清洗策略")
    with store.connect() as c:
        row = c.execute(
            "SELECT v.* FROM article_versions v JOIN articles a ON a.current_version=v.id WHERE a.id=?",
            (article_pk,),
        ).fetchone()
        if not row:
            raise ValueError("文章不存在")
        content = (
            clean(row["raw_content"]) if mode == "preserve-v1" else row["raw_content"]
        )
        if mode == "html-readable-v1":
            from .content import readable_html

            content = readable_html(row["raw_content"])
        if row["clean_version"] == mode:
            return row["id"]
        digest = hashlib.sha256(
            dumps(
                [row["title"], row["raw_content"], content, mode, row["metadata_json"]]
            ).encode()
        ).hexdigest()
        fields = [
            "article_fk",
            "title",
            "raw_content",
            "clean_content",
            "content_hash",
            "metadata_json",
            "parse_status",
            "clean_version",
            "source_updated_at",
            "import_batch_id",
            "source_file",
            "created_at",
        ]
        values = [
            row["article_fk"],
            row["title"],
            row["raw_content"],
            content,
            digest,
            row["metadata_json"],
            parse_status(row["title"], content),
            mode,
            row["source_updated_at"],
            row["import_batch_id"],
            row["source_file"],
            time.time(),
        ]
        version = c.execute(
            "INSERT INTO article_versions("
            + ",".join(fields)
            + ") VALUES("
            + ",".join("?" for _ in fields)
            + ")",
            values,
        ).lastrowid
        c.execute(
            "UPDATE articles SET current_version=? WHERE id=?", (version, article_pk)
        )
    store.audit(
        "clean_strategy", article_pk, "新版本采用 " + mode + "；原文与历史结果不变"
    )
    return version
