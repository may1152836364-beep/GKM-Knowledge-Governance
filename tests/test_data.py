from openpyxl import Workbook, load_workbook
from gkm.storage import Store
from gkm.excel import import_xlsx, preview, safe_cell


def fixture(path, body=None):
    w = Workbook()
    s = w.active
    s.append(
        [
            "Article ID(文章ID)",
            "title(文章标题)",
            "content(文章正文)",
            "Language(语言)",
            "未知列",
        ]
    )
    s.append(
        [
            "00001",
            "维修说明",
            body
            or "这是仍然有用的维修流程。取消标记和删除是正常的操作步骤。请先联系服务团队。",
            "zh",
            "附加信息",
        ]
    )
    s.append(
        [
            "00002",
            "Repair guide",
            "Repair procedures remain useful for owners of older devices. Please contact the service team.",
            "en",
            "metadata",
        ]
    )
    s.append(["", "无 ID", "正文", "zh", ""])
    w.save(path)


def test_import_versions(tmp_path):
    path = tmp_path / "input.xlsx"
    fixture(path)
    db = Store(tmp_path / "data")
    p = preview(path)
    r = import_xlsx(db, path, p["mapping"])
    assert (r["inserted"], r["skipped"]) == (2, 1)
    assert db.articles()[0]["article_id"] == "00001"
    v = db.article_version(db.articles()[0]["current_version"])
    assert "取消标记和删除" in v["clean_content"] and "未知列" in v["metadata_json"]
    assert import_xlsx(db, path, p["mapping"])["duplicates"] == 2
    fixture(
        path,
        "这是更新后的业务流程，完整内容必须保留并新建版本，不能覆盖原来的正文及分析记录。",
    )
    assert import_xlsx(db, path, p["mapping"])["updated"] == 1
    assert db.one("SELECT count(*) n FROM article_versions")["n"] == 3
    assert len(db.articles(query="Repair", page_size=1)) == 1
    db.backup(tmp_path / "backup.sqlite3")
    db.restore(tmp_path / "backup.sqlite3")
    assert db.one("SELECT count(*) n FROM articles")["n"] == 2


def test_export_formula():
    assert safe_cell('=HYPERLINK("x")').startswith("'")
    assert safe_cell("  @attack").startswith("'")
    assert safe_cell("中文") == "中文"


def test_streaming_sheet_without_dimensions_and_reversion(tmp_path):
    path = tmp_path / "stream.xlsx"
    w = Workbook(write_only=True)
    s = w.create_sheet("sheet")
    s.append(["文章ID", "文章标题", "文章正文", "语言"])
    s.append(
        [
            "001",
            "原始",
            "原始业务步骤和必要说明，包含足够的业务信息供模型逐段判断并且不能随意删除。",
            "zh",
        ]
    )
    w.save(path)
    p = preview(path)
    db = Store(tmp_path / "db")
    assert import_xlsx(db, path, p["mapping"])["inserted"] == 1
    old = db.article_version(1)["raw_content"]
    w = load_workbook(path)
    w.active["C2"] = (
        "更新后的正文仍然提供有效业务信息，需要新增历史版本，不能覆盖之前已经存在的记录。"
    )
    w.save(path)
    w.close()
    assert import_xlsx(db, path, p["mapping"])["updated"] == 1
    w = load_workbook(path)
    w.active["C2"] = old
    w.save(path)
    w.close()
    assert import_xlsx(db, path, p["mapping"])["updated"] == 1
    assert db.one("SELECT count(*) n FROM article_versions")["n"] == 3


def test_schema_migration_keeps_history_and_rolls_back_unknown_version(tmp_path):
    import sqlite3
    import pytest
    from gkm.storage import SCHEMA_SQL

    directory = tmp_path / "old"
    directory.mkdir()
    path = directory / "gkm.sqlite3"
    old = SCHEMA_SQL.replace(
        "created_at REAL NOT NULL);",
        "created_at REAL NOT NULL, UNIQUE(article_fk,content_hash));",
    ).replace(
        ", case_prompt TEXT NOT NULL DEFAULT '参考案例仅用于相同事实和适用边界；不可把案例当作当前文章外部事实'",
        "",
    )
    with sqlite3.connect(path) as c:
        c.executescript(old)
        c.execute("PRAGMA user_version=1")
    db = Store(directory)
    assert db.one("PRAGMA user_version")["user_version"] == 2
    assert list(directory.glob("before-migration-*.sqlite3"))
    assert "case_prompt" in db.one("SELECT * FROM prompt_versions")
    assert db.one("PRAGMA foreign_key_check") is None
    db.execute("PRAGMA user_version=99")
    with pytest.raises(RuntimeError):
        Store(directory)


def test_cleaning_undo_preserves_source_and_history(tmp_path):
    from gkm.excel import set_cleaning

    path = tmp_path / "data.xlsx"
    fixture(
        path,
        "业务操作\r\n取消标记和删除是业务步骤，历史知识具有使用价值。请勿删除原始正文。",
    )
    db = Store(tmp_path / "db")
    import_xlsx(db, path, preview(path)["mapping"])
    row = db.articles()[0]
    old = db.article_version(row["current_version"])
    new = set_cleaning(db, row["id"], "raw-v1")
    current = db.article_version(new)
    assert (
        current["clean_content"] == old["raw_content"]
        and current["raw_content"] == old["raw_content"]
    )
    assert (
        new != old["id"]
        and db.article_version(old["id"])["clean_version"] == "preserve-v1"
    )


def test_excel_shared_strings_spill_to_disk(tmp_path):
    import zipfile
    from xml.etree import ElementTree as ET
    from gkm.xlsx_stream import Workbook as StreamingWorkbook

    path = tmp_path / "inline.xlsx"
    fixture(path)
    target = tmp_path / "shared.xlsx"
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    strings = []
    root = ET.Element(ns + "sst")
    with (
        zipfile.ZipFile(path) as src,
        zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for name in src.namelist():
            data = src.read(name)
            if name == "xl/worksheets/sheet1.xml":
                tree = ET.fromstring(data)
                for cell in tree.iter(ns + "c"):
                    if cell.get("t") == "inlineStr":
                        value = "".join(e.text or "" for e in cell.iter(ns + "t"))
                        inline = cell.find(ns + "is")
                        if inline is not None:
                            cell.remove(inline)
                        cell.set("t", "s")
                        ET.SubElement(cell, ns + "v").text = str(len(strings))
                        strings.append(value)
                        si = ET.SubElement(root, ns + "si")
                        ET.SubElement(si, ns + "t").text = value
                data = ET.tostring(tree)
            elif name == "[Content_Types].xml":
                tree = ET.fromstring(data)
                ET.SubElement(
                    tree,
                    "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
                    PartName="/xl/sharedStrings.xml",
                    ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
                )
                data = ET.tostring(tree)
            dst.writestr(name, data)
        dst.writestr("xl/sharedStrings.xml", ET.tostring(root))
    wb = StreamingWorkbook(target)
    cache = wb.reader.disk_strings.temp.name
    assert next(wb[wb.sheetnames[0]].iter_rows())[0].value == "Article ID(文章ID)"
    wb.close()
    from pathlib import Path

    assert not Path(cache).exists()
    db = Store(tmp_path / "data")
    r = import_xlsx(db, target, preview(target)["mapping"])
    assert r["inserted"] == 2 and db.articles()[0]["article_id"] == "00001"


def test_insufficient_input_and_formula_isolation(tmp_path):
    from gkm.excel import parse_status

    assert (
        parse_status("标题", "https://video.example.com/123") == "insufficient_content"
    )
    assert parse_status("", "完整正文") == "missing_title"
    assert parse_status("标题", "") == "missing_body"
    path = tmp_path / "formula.xlsx"
    fixture(path)
    w = load_workbook(path)
    w.active["C2"] = '=HYPERLINK("https://example.com")'
    w.save(path)
    w.close()
    db = Store(tmp_path / "db")
    report = import_xlsx(db, path, preview(path)["mapping"])
    assert report["inserted"] == 1 and report["skipped"] == 2
    assert all(row["article_id"] != "00001" for row in db.articles())


def test_html_is_readable_without_executing_or_losing_source(tmp_path):
    from gkm.excel import set_cleaning

    raw = '<h1>维修流程</h1><table><tr><td>取消标记</td><td>删除操作</td></tr></table><a href="https://approved.example/help">售后链接</a><img src="https://example/image.png" alt="拆卸图示"><script>alert("不执行")</script>'
    path = tmp_path / "html.xlsx"
    fixture(path, raw)
    db = Store(tmp_path / "db")
    import_xlsx(db, path, preview(path)["mapping"])
    row = db.articles()[0]
    vid = set_cleaning(db, row["id"], "html-readable-v1")
    v = db.article_version(vid)
    assert v["raw_content"] == raw and "<table>" not in v["clean_content"]
    for value in (
        "取消标记",
        "删除操作",
        "https://approved.example/help",
        "拆卸图示",
        "https://example/image.png",
    ):
        assert value in v["clean_content"]
    assert set_cleaning(db, row["id"], "raw-v1") != vid
